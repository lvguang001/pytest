# -*- coding: utf-8 -*-
"""文书工具箱的单元测试（`documents.py`）。

docx 相关的用例自己现造 fixture（python-docx 已经在依赖里），不碰真实模板、
不碰网络、不碰界面。
"""

import os

import pytest
from docx import Document

import documents
from documents import (
    build_notice_template_data, insert_questions_to_document,
    normalize_compact_date, notice_file_name, notice_template_name,
    read_approval_table, tick_approval_box, unique_path,
)


def _docx_with_table(path, rows):
    """造一份只含一张表格的 docx：rows 是 [[单元格文本, ...], ...]"""
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    doc.save(str(path))
    return str(path)


# ============================================================================
# 纯格式转换
# ============================================================================

@pytest.mark.parametrize("raw,expected", [
    ("20260720", "2026年07月20日"),
    ("20260101", "2026年01月01日"),
    ("2026年7月20日", "2026年7月20日"),   # 已是中文，原样
    ("2026072", "2026072"),               # 位数不对
    ("2026072a", "2026072a"),             # 含非数字
    ("", ""),
    (None, None),
])
def test_normalize_compact_date(raw, expected):
    assert normalize_compact_date(raw) == expected


def test_notice_template_and_file_name_follow_the_conclusion():
    assert notice_template_name("不予认定") == "不予工伤认定告知书（样本）.docx"
    assert notice_template_name("予以认定") == "工伤认定告知书（样本）.docx"
    assert notice_template_name("") == "工伤认定告知书（样本）.docx"

    assert notice_file_name("不予认定", "张三") == "张三不予工伤认定告知书.docx"
    assert notice_file_name("予以认定", "张三") == "张三工伤认定告知书.docx"


def test_unique_path_avoids_clobbering(tmp_path):
    """审批表每生成一次就多一份，不能覆盖上一份。"""
    d = str(tmp_path)
    assert unique_path(d, "张三案件审批表").endswith("张三案件审批表.docx")
    open(unique_path(d, "张三案件审批表"), "w").close()          # 占用第一个
    assert unique_path(d, "张三案件审批表").endswith("张三案件审批表(2).docx")
    open(unique_path(d, "张三案件审批表"), "w").close()          # 再占用第二个
    assert unique_path(d, "张三案件审批表").endswith("张三案件审批表(3).docx")


# ============================================================================
# 模板数据
# ============================================================================

def test_build_notice_template_data_maps_the_case_object():
    case = {"case_id": "案本2026001", "name": "张三", "id_card": "3301",
            "labor_unit": "某公司", "injury_description": "砸伤",
            "medical_conclusion": "骨折", "unit_type": "企业", "identity": "职工"}
    d = build_notice_template_data(case)
    assert d["本人姓名"] == "张三"
    assert d["用人单位"] == "某公司"
    assert d["受伤经过"] == "砸伤"
    assert d["案本号"] == "案本2026001" == d["受理编号"]
    assert "认定依据句" in d and "本人所属表述" in d


def test_build_notice_template_data_falls_back_when_fields_are_missing():
    d = build_notice_template_data({})
    assert d["受伤经过"] == "详见谈话笔录"     # 两个来源都没有时
    assert d["本人姓名"] == ""
    assert d["申请时间"] and d["受理时间"]     # 补成今天


def test_notice_basis_sentence_switches_on_the_conclusion():
    case = {"proposed_article": "第十四条第（一）项", "unit_type": "企业"}
    yes = build_notice_template_data({**case, "conclusion": "予以认定"})["认定依据句"]
    no = build_notice_template_data({**case, "conclusion": "不予认定"})["认定依据句"]
    assert "认定为工伤" in yes and "不予" not in yes
    assert "不予认定" in no


# ============================================================================
# 读审批表
# ============================================================================

def test_read_approval_table_takes_the_cell_to_the_right(tmp_path):
    path = _docx_with_table(tmp_path / "审批表.docx", [
        ["用人单位", "某某建筑劳务有限公司"],
        ["职工姓名", "张三"],
        ["身份证号", "330102199001011234"],
    ])
    data = read_approval_table(path)
    assert data == {
        "用人单位": "某某建筑劳务有限公司",
        "职工姓名": "张三",
        "职工身份证号": "330102199001011234",
    }


def test_read_approval_table_maps_the_medical_column(tmp_path):
    """搜的是「医疗诊断」，存进模板的字段名却是「医疗证明」。"""
    path = _docx_with_table(tmp_path / "审批表.docx", [["医疗诊断", "右足跖骨骨折"]])
    assert read_approval_table(path)["医疗证明"] == "右足跖骨骨折"


def test_read_approval_table_ignores_empty_and_self_repeating_cells(tmp_path):
    """右边那格为空、或又写了一遍关键词，都不算提取到。"""
    path = _docx_with_table(tmp_path / "审批表.docx", [
        ["用人单位", ""],
        ["职工姓名", "职工姓名"],
    ])
    assert read_approval_table(path) == {}


def test_read_approval_table_ignores_a_keyword_in_the_last_column(tmp_path):
    """关键词在最后一格时右边没东西可取，不能越界。"""
    path = _docx_with_table(tmp_path / "审批表.docx", [["张三", "用人单位"]])
    assert read_approval_table(path) == {}


def test_read_approval_table_keeps_the_first_hit(tmp_path):
    path = _docx_with_table(tmp_path / "审批表.docx", [
        ["用人单位", "第一家"],
        ["用人单位", "第二家"],
    ])
    assert read_approval_table(path)["用人单位"] == "第一家"


# ============================================================================
# 勾选审批表的结论框
# ============================================================================

def _cell_texts(path):
    return [c.text for t in Document(path).tables for r in t.rows for c in r.cells]


def test_tick_approval_box_checks_the_confirm_box(tmp_path):
    src = _docx_with_table(tmp_path / "t.docx", [["□认定工伤    □不予认定工伤"]])
    out = tick_approval_box(src, "予以认定")
    try:
        assert "☑认定工伤" in _cell_texts(out)[0]
        assert "□不予认定工伤" in _cell_texts(out)[0], "另一格不该被动"
    finally:
        os.remove(out)


def test_tick_approval_box_checks_the_deny_box(tmp_path):
    src = _docx_with_table(tmp_path / "t.docx", [["□认定工伤    □不予认定工伤"]])
    out = tick_approval_box(src, "不予认定")
    try:
        assert "☑不予认定工伤" in _cell_texts(out)[0]
        assert "□认定工伤" in _cell_texts(out)[0]
    finally:
        os.remove(out)


def test_tick_approval_box_never_touches_the_original_template(tmp_path):
    """模板是共用的，勾选只能改临时副本。"""
    src = _docx_with_table(tmp_path / "t.docx", [["□认定工伤"]])
    before = _cell_texts(src)
    out = tick_approval_box(src, "予以认定")
    try:
        assert _cell_texts(src) == before
        assert out != src
    finally:
        os.remove(out)


# ============================================================================
# 往笔录末尾追加问题
# ============================================================================

def test_insert_questions_appends_all_of_them(tmp_path):
    path = _docx_with_table(tmp_path / "笔录.docx", [["问：姓名？", "答：张三"]])
    ok, msg = insert_questions_to_document(path, ["是否签合同？", "谁在现场？"])
    assert ok and msg == "插入成功"
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "AI建议补充问题：" in text
    assert "1. 是否签合同？" in text
    assert "2. 谁在现场？" in text


def test_insert_questions_reports_failure_instead_of_raising(tmp_path):
    """文件不存在时返回 (False, 原因)，让调用方去弹窗。"""
    ok, msg = insert_questions_to_document(str(tmp_path / "没有这个.docx"), ["x"])
    assert ok is False
    assert msg
