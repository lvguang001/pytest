# -*- coding: utf-8 -*-
"""一次性补丁：谈话模板头部 + 告知书正文，支持 单位性质/身份 措辞（机关公务员/事业单位）。
运行前已备份原 docx 为同路径 *.bak.docx。可安全重复运行（幂等）。"""
import glob, os, shutil
from docx import Document

TALK_DIR = 'resource/模板文件/谈话模板'
DOC_DIR = 'resource/模板文件/文书模板'

ROLE_TEXT = {
    '本人': '工作单位： {{单位名称}}  职务或岗位： {{本人岗位}}  单位性质：{{单位性质}}  身份：{{本人身份}}',
    '证人': '工作单位： {{单位名称}}  职务或岗位： {{证人岗位}}  单位性质：{{单位性质}}  身份：{{证人身份}}',
    '法人': '工作单位： {{单位名称}}  职务或岗位： {{法人职务}}  单位性质：{{单位性质}}  身份：{{法人身份}}',
    '家属': '工作单位： {{单位名称}}  职务或岗位： {{与死者关系}}  单位性质：{{单位性质}}  身份：{{家属身份}}',
}


def backup(path):
    bak = path + '.bak.docx'
    if not os.path.exists(bak):
        shutil.copy2(path, bak)


def set_para_text(p, text):
    """用第一 run 覆盖整段（保留其字体/格式），删除其余 run。"""
    if not p.runs:
        p.add_run(text)
        return
    first = p.runs[0]
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    first.text = text


def patch_talk():
    changed = []
    for role, target in ROLE_TEXT.items():
        for f in glob.glob(os.path.join(TALK_DIR, f'*{role}谈话笔录*.docx')):
            if f.endswith('.bak.docx'):
                continue
            backup(f)
            doc = Document(f)
            hit = False
            for p in doc.paragraphs:
                if '工作单位' in p.text and '{{' in p.text:
                    if p.text.strip() != target.strip():
                        set_para_text(p, target)
                        hit = True
            if hit:
                doc.save(f)
                changed.append(os.path.basename(f))
    return changed


def patch_notices():
    changed = []
    old_subject = '系{{用人单位}}职工，从事金工工作'
    old_basis = '符合《工伤保险条例》第十四条第（一）项认定工伤之规定。现拟决定认定为工伤。'
    for f in glob.glob(os.path.join(DOC_DIR, '*告知书（样本）.docx')):
        if f.endswith('.bak.docx'):
            continue
        backup(f)
        doc = Document(f)
        hit = False
        for p in doc.paragraphs:
            t = p.text
            if old_subject in t or old_basis in t:
                n = t.replace(old_subject, '{{本人所属表述}}').replace(old_basis, '{{认定依据句}}')
                set_para_text(p, n)
                hit = True
        if hit:
            doc.save(f)
            changed.append(os.path.basename(f))
    return changed


if __name__ == '__main__':
    print('谈话模板补丁:', patch_talk())
    print('告知书补丁:', patch_notices())
