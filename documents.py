# -*- coding: utf-8 -*-
"""文书生成工具箱：读审批表、建模板数据、渲染 docx 模板。

原先这些都散在 MainWindow 的方法里（2026-09 抽出来），和界面反馈
（状态栏、弹窗、打开文件）缠在一起。这里只留**不依赖界面**的那半：
输入是案件数据或文件路径，输出是字典或文件——因此可以直接单测。

编排（取哪个目录、弹什么窗、更新送达流程）仍留在 MainWindow。
"""

import logging
import os
import shutil
import tempfile

from docx import Document
from docxtpl import DocxTemplate

from case_classifier import (
    DEFAULT_IDENTITY, DEFAULT_UNIT_TYPE, notice_basis_sentence, person_affiliation,
)
from services import date_now, time_now

logger = logging.getLogger(__name__)


# ============================================================================
# 读审批表
# ============================================================================

# 搜索词 → 提取出来的字段名。注意「医疗诊断」搜的是「医疗证明」这个字段。
APPROVAL_FIELD_MAP = {
    '用人单位': '用人单位',
    '职工姓名': '职工姓名',
    '身份证号': '职工身份证号',
    '申请时间': '申请时间',
    '受理时间': '受理时间',
    '受伤经过': '受伤经过',
    '医疗诊断': '医疗证明',  # 搜索"医疗诊断"
}

# 生成告知书/通知书时缺一不可的字段
APPROVAL_REQUIRED_FIELDS = ['用人单位', '职工姓名', '职工身份证号',
                            '申请时间', '受理时间', '受伤经过', '医疗证明']


def read_approval_table(file_path: str) -> dict:
    """从审批表 docx 的表格里提取数据：找到关键词单元格，取它右边那格。

    只做「读」，找不到的字段就不放进结果——补默认值是调用方的事
    （它得看界面上的值）。
    """
    print(f"📄 开始提取审批表数据: {file_path}")
    extracted_data = {}
    document = Document(file_path)

    for table in document.tables:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell_text = cell.text.strip()

                for search_term, data_field in APPROVAL_FIELD_MAP.items():
                    if data_field in extracted_data:      # 已找到，跳过
                        continue

                    if search_term in cell_text:
                        print(f"✅ 找到关键词 '{search_term}'")

                        if i + 1 < len(row.cells):
                            right_text = row.cells[i + 1].text.strip()
                            # 右边那格得有内容、且不是关键词本身
                            if right_text and right_text != search_term:
                                extracted_data[data_field] = right_text
                                print(f"  提取 {data_field}: {right_text}")
                            else:
                                logger.warning(f"  ⚠️ {data_field}: 右边单元格为空")
                        else:
                            logger.warning(f"  ⚠️ {data_field}: 没有右侧单元格")

    print("\n📋 提取结果:")
    for field in APPROVAL_REQUIRED_FIELDS:
        if field in extracted_data:
            print(f"  ✅ {field}: {extracted_data[field]}")
        else:
            logger.error(f"  ❌ {field}: 未找到")

    return extracted_data


# ============================================================================
# 读任意 docx 的正文
# ============================================================================

def read_docx_text(file_path: str) -> str:
    """读 docx 的**段落**正文（不含表格），每段一行、空段丢掉；读不了返回空串。

    只读段落是有意的：谈话笔录的问答都在段落里，表格是审批表那种才有的东西。
    读不出内容不抛异常——调用方（如「有没有本人笔录」的判断）要把它当成「没有」，
    而不是让一个坏文件把整条生成链打断。
    """
    try:
        doc = Document(file_path)
    except Exception as e:
        logger.warning("⚠️ 读取 docx 失败 %s: %s", file_path, e)
        return ""
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def normalize_compact_date(value) -> str:
    """8 位数字的日期（20260720）补成「2026年07月20日」；其余原样返回。

    审批表里手填的日期常常是紧凑格式，直接塞进模板会很难看。
    """
    if isinstance(value, str) and value.isdigit() and len(value) == 8:
        return f"{value[0:4]}年{value[4:6]}月{value[6:8]}日"
    return value


# ============================================================================
# 模板数据（纯：案件数据 → 模板占位符字典）
# ============================================================================

def build_notice_template_data(case_obj: dict) -> dict:
    """构建工伤告知书模板字典（从案件 JSON 数据）"""
    current_date = date_now()
    return {
        '本人姓名': case_obj.get('name', ''),
        '本人身份证号': case_obj.get('id_card', ''),
        '用人单位': case_obj.get('labor_unit', ''),  # 用人单位（签合同单位）
        '受伤经过': case_obj.get('injury_process', case_obj.get('injury_description', '详见谈话笔录')),
        '医疗证明': case_obj.get('medical_conclusion', ''),
        '申请时间': case_obj.get('apply_time', current_date),
        '受理时间': case_obj.get('accept_time', current_date),
        '当前时期': current_date + time_now(),
        '告知日期': current_date,
        '受理编号': case_obj.get('case_id', ''),
        '案本号': case_obj.get('case_id', ''),
        '单位性质': case_obj.get('unit_type', DEFAULT_UNIT_TYPE),
        '本人身份': case_obj.get('identity', DEFAULT_IDENTITY),
        '本人所属表述': person_affiliation(case_obj),
        '认定依据句': notice_basis_sentence(
            case_obj, deny=('不予' in str(case_obj.get('conclusion', '')))),
    }


def notice_template_name(conclusion: str) -> str:
    """按结论选模板：不予认定用「不予」那份。"""
    return ('不予工伤认定告知书（样本）.docx' if conclusion == "不予认定"
            else '工伤认定告知书（样本）.docx')


def notice_file_name(conclusion: str, person_name: str) -> str:
    return (f"{person_name}不予工伤认定告知书.docx" if conclusion == "不予认定"
            else f"{person_name}工伤认定告知书.docx")


# ============================================================================
# 渲染 docx 模板
# ============================================================================

def render_template(template_path: str, data: dict) -> DocxTemplate:
    """渲染一份 docx 模板，返回**尚未保存**的 DocxTemplate。

    不给路径是因为调用方要自己决定存哪（审批表还要先做一次勾选预处理、
    并且要避开同名文件）。
    """
    word = DocxTemplate(template_path)
    word.render(data)
    return word


def unique_path(folder: str, base_name: str, ext: str = ".docx") -> str:
    """在 folder 下取一个不重名的路径：<名>.docx / <名>(2).docx / <名>(3).docx …

    审批表每生成一次就多一份，不能覆盖上一份。
    """
    target = os.path.join(folder, base_name + ext)
    counter = 2
    while os.path.exists(target):
        target = os.path.join(folder, f"{base_name}({counter}){ext}")
        counter += 1
    return target


# ============================================================================
# 审批表：按结论勾选「认定工伤 / 不予认定工伤」
# ============================================================================

def tick_approval_box(template_path: str, conclusion: str) -> str:
    """把模板里的 `□认定工伤` / `□不予认定工伤` 按结论改成 `☑`，返回新文件路径。

    原模板不能动，所以先复制到临时文件再改；调用方用完自己删。
    """
    temp_template = os.path.join(tempfile.gettempdir(), '_temp_approval_table.docx')
    shutil.copy2(template_path, temp_template)

    doc_edit = Document(temp_template)
    check_confirm = (conclusion == "予以认定")
    for table in doc_edit.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if check_confirm and "□认定工伤" in p.text:
                        full = p.text.replace("□认定工伤", "☑认定工伤", 1)
                    elif (not check_confirm) and "□不予认定工伤" in p.text:
                        full = p.text.replace("□不予认定工伤", "☑不予认定工伤", 1)
                    else:
                        continue
                    # 整段文字塞进第一个 run，其余清空（保留首个 run 的字体）
                    if p.runs:
                        p.runs[0].text = full
                        for r in p.runs[1:]:
                            r.text = ""
                    break
    doc_edit.save(temp_template)
    return temp_template


# ============================================================================
# 往笔录 docx 末尾追加补充问题
# ============================================================================

def insert_questions_to_document(file_path: str, questions: list) -> tuple:
    """将问题插入到Word文档中

    Args:
        file_path: Word文档路径
        questions: 问题列表

    Returns:
        (是否成功, 消息)
    """
    try:
        from docx.shared import RGBColor

        print(f"📄 正在处理文档: {file_path}")
        doc = Document(file_path)

        # 添加分隔线
        separator = doc.add_paragraph("=" * 50)
        separator.alignment = 1  # 居中

        # 添加标题
        title = doc.add_paragraph("AI建议补充问题：")
        title.runs[0].bold = True
        title.alignment = 0  # 左对齐

        # 插入每个问题 + 答案占位符
        for i, question_text in enumerate(questions, 1):
            question_para = doc.add_paragraph()
            question_para.add_run(f"{i}. {question_text}")

            answer_para = doc.add_paragraph()
            answer_run = answer_para.add_run("答：")
            answer_run.font.underline = True
            answer_run.font.color.rgb = RGBColor(0, 0, 0)

            doc.add_paragraph()      # 空行

        doc.save(file_path)
        print(f"✅ 成功插入 {len(questions)} 个问题到文档")
        return True, "插入成功"

    except Exception as e:
        logger.error(f"❌ 插入问题到文档失败: {e}")
        import traceback
        traceback.print_exc()
        return False, str(e)
